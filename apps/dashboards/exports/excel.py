"""
Excel report builder (openpyxl).

Produces a styled .xlsx workbook of room utilisation by building, suitable for
planning meetings (concept note 7.5). Returns a Django HttpResponse ready to
stream to the browser.
"""
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone


def build_utilisation_xlsx(by_building: list[dict], title="Room Utilisation by Building") -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Utilisation"

    # Title row
    ws.merge_cells("A1:E1")
    cell = ws["A1"]
    cell.value = f"University of Zimbabwe — {title}"
    cell.font = Font(size=13, bold=True, color="0E1B3D")
    cell.alignment = Alignment(horizontal="left")

    ws["A2"] = f"Generated: {timezone.now():%d %b %Y %H:%M}"
    ws["A2"].font = Font(size=9, italic=True, color="666666")

    # Header row
    headers = ["Building", "Rooms", "Total Capacity", "In Use / Booked", "Utilisation %"]
    header_fill = PatternFill("solid", fgColor="0E1B3D")
    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill

    # Data rows
    row = 5
    for r in by_building:
        ws.cell(row=row, column=1, value=r["label"])
        ws.cell(row=row, column=2, value=r["rooms"])
        ws.cell(row=row, column=3, value=r["capacity"])
        ws.cell(row=row, column=4, value=r["used"])
        ws.cell(row=row, column=5, value=r["utilisation"] / 100).number_format = "0.0%"
        row += 1

    # Column widths
    for col, width in zip("ABCDE", [34, 10, 16, 16, 14]):
        ws.column_dimensions[col].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    stamp = timezone.now().strftime("%Y%m%d")
    response["Content-Disposition"] = f'attachment; filename="uz-room-utilisation-{stamp}.xlsx"'
    return response
